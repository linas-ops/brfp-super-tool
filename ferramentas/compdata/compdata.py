# -*- coding: utf-8 -*-
"""
Modelo de dominio do compdata do FC.

O ponto central: em compobj.txt o ID do no E o numero da linha. Qualquer
insercao, remocao ou reordenacao renumera tudo que vem depois, e cada
referencia a no espalhada pelos outros arquivos precisa ser remapeada.
Toda mutacao aqui passa por rebuild(), que faz isso de uma vez so.
"""
import io
import os
import re
import copy
import shutil
import datetime
import collections

# ---------------------------------------------------------------- constantes

NODE_TYPES = {
    -1: 'Reservado',
    0: 'Mundo',
    1: 'Confederacao',
    2: 'Pais',
    3: 'Competicao',
    4: 'Estagio',
    5: 'Grupo',
    6: 'Competicao especial',
}

CONTAINER_TYPES = (0, 1, 2)
COMP_TYPES = (3, 6)

# settings cujo VALOR e um ID de no. Verificado uma a uma contra o arquivo real:
# tudo que ficou de fora (info_special_team_id, nation_id, match_stadium,
# info_league_promo/releg, schedule_refcomp, asset_id...) e ID de banco ou de
# asset -- remapear corromperia os dados.
SETTING_NODE_KEYS = {
    'rule_suspension', 'schedule_forcecomp', 'advance_pointskeep',
    'standings_checkrank', 'advance_standingskeep', 'advance_standingsrank',
    'advance_maxteamsstageref', 'schedule_use_dates_comp',
    'advance_teamcompdependency', 'advance_2026_wc_third_place_setup',
}

# posicao (indice do campo) que e ID de no, por verbo de task.
# formato da linha: comp, quando, verbo, f3, f4, f5, f6
TASK_NODE_FIELDS = {
    'FillWithTeam': [3],
    'FillFromSpecialTeams': [3],
    'FillFromSpecialTeamsWithNation': [3],
    'FillFromLeague': [3],
    'FillFromLeagueInOrder': [3],
    'FillFromLeagueMaxFromCountry': [3],
    'FillFromTopCoefficientCountry': [3],
    'FillFromTopYear5CoefficientCountry': [3],
    'ClearLeagueStats': [3],
    'UpdateLeagueStats': [3],
    'UpdateMultiGroupLeagueStats': [3],
    'FillFromCompTable': [3, 4],
    'FillFromCompTableBackupLeague': [3, 4],
    'FillFromCompTablePosBackupSameLeague': [3, 4],
    'FillFromCompTableBackup': [3, 4, 5],
    'UpdateTable': [3, 4],
}

STAGE_TYPES = ['SETUP', 'LEAGUE', 'KO1LEG', 'KO2LEGS']
MATCH_SITUATIONS = ['LEAGUE', 'GROUP', 'QUALIFY', 'ROUNDX', 'QUARTER', 'SEMI',
                    'THIRDPLACE', 'FINAL']
COMP_KINDS = ['LEAGUE', 'CUP', 'SUPERCUP', 'PLAYOFF', 'INTERCUP', 'INTERQUAL', 'NONE']
SORT_KEYS = ['POINTS', 'WINS', 'GOALDIFF', 'GOALSFOR', 'H2HPOINTS', 'H2HGOALDIFF',
             'H2HGOALSFOR', 'TEAMRATING', 'PREVRANK']

# arquivos com uma unica coluna de no (coluna 0)
SIMPLE_FILES = ['standings', 'schedule', 'objectives', 'initteams', 'weather']


def _read(path):
    if not os.path.exists(path):
        return []
    with io.open(path, 'r', encoding='utf-8', newline='') as f:
        raw = f.read()
    return [l for l in raw.replace('\r\n', '\n').split('\n') if l.strip()]


def _write(path, lines):
    with io.open(path, 'w', encoding='utf-8', newline='') as f:
        for l in lines:
            f.write(l + '\r\n')


class Compdata(object):
    """Carrega, edita e grava o conjunto de arquivos do compdata."""

    def __init__(self, path):
        self.path = path
        self.load()

    # ------------------------------------------------------------------ io

    def load(self):
        p = lambda n: os.path.join(self.path, n + '.txt')

        self.nodes = []
        for line in _read(p('compobj')):
            f = line.split(',')
            self.nodes.append({'type': int(f[1]), 'short': f[2],
                               'desc': f[3], 'parent': int(f[4])})

        self.settings = [l.split(',', 2) for l in _read(p('settings'))]
        self.settings = [[int(a), b, c] for a, b, c in self.settings]

        self.rows = {}
        for name in SIMPLE_FILES:
            self.rows[name] = [l.split(',') for l in _read(p(name))]

        self.advancement = [[int(x) for x in l.split(',')]
                            for l in _read(p('advancement'))]
        self.tasks = [l.split(',') for l in _read(p('tasks'))]

        self._history = []
        self._future = []
        self._batch = 0
        self.dirty = False
        self.baseline = self.line_counts()

    def save(self, backup=True):
        stamp = datetime.datetime.now().strftime('%Y%m%d_%H%M%S')
        bdir = os.path.join(self.path, '_backup_editor_' + stamp)
        names = ['compobj', 'compids', 'settings', 'advancement', 'tasks'] + SIMPLE_FILES
        if backup:
            os.makedirs(bdir)
            for n in names:
                src = os.path.join(self.path, n + '.txt')
                if os.path.exists(src):
                    shutil.copy2(src, bdir)

        p = lambda n: os.path.join(self.path, n + '.txt')
        _write(p('compobj'), ['%d,%d,%s,%s,%d' % (i, n['type'], n['short'],
                                                  n['desc'], n['parent'])
                              for i, n in enumerate(self.nodes)])
        _write(p('compids'), [str(i) for i, n in enumerate(self.nodes)
                              if n['type'] in COMP_TYPES])
        _write(p('settings'), ['%d,%s,%s' % tuple(r) for r in self.settings])
        _write(p('advancement'), [','.join(str(x) for x in r) for r in self.advancement])
        _write(p('tasks'), [','.join(r) for r in self.tasks])
        for name in SIMPLE_FILES:
            _write(p(name), [','.join(r) for r in self.rows[name]])

        self.dirty = False
        return bdir if backup else None

    def line_counts(self):
        out = {
            'compobj': len(self.nodes),
            'compids': sum(1 for n in self.nodes if n['type'] in COMP_TYPES),
            'settings': len(self.settings),
            'advancement': len(self.advancement),
            'tasks': len(self.tasks),
        }
        for name in SIMPLE_FILES:
            out[name] = len(self.rows[name])
        return out

    # ------------------------------------------------------------- historico

    def snapshot(self):
        """Um passo de desfazer. Dentro de um batch, so o primeiro conta --
        assim o assistente de competicao desfaz de uma vez, nao no de no."""
        if self._batch:
            self._batch += 1
            return
        self._history.append(copy.deepcopy(
            (self.nodes, self.settings, self.rows, self.advancement, self.tasks)))
        if len(self._history) > 60:
            self._history.pop(0)
        self._future = []
        self.dirty = True

    def begin_batch(self):
        self.snapshot()
        self._batch = 1

    def end_batch(self):
        self._batch = 0

    def _restore(self, snap):
        self.nodes, self.settings, self.rows, self.advancement, self.tasks = \
            copy.deepcopy(snap)

    def undo(self):
        if not self._history:
            return False
        self._future.append(copy.deepcopy(
            (self.nodes, self.settings, self.rows, self.advancement, self.tasks)))
        self._restore(self._history.pop())
        self.dirty = True
        return True

    def redo(self):
        if not self._future:
            return False
        self._history.append(copy.deepcopy(
            (self.nodes, self.settings, self.rows, self.advancement, self.tasks)))
        self._restore(self._future.pop())
        self.dirty = True
        return True

    # ---------------------------------------------------------------- arvore

    def children(self, nid):
        return [i for i, n in enumerate(self.nodes) if n['parent'] == nid]

    def subtree(self, nid):
        out = [nid]
        for k in self.children(nid):
            out += self.subtree(k)
        return out

    def ancestors(self, nid):
        out = []
        seen = 0
        while nid is not None and self.nodes[nid]['parent'] != -1 and seen < 60:
            nid = self.nodes[nid]['parent']
            out.append(nid)
            seen += 1
        return out

    def owning_comp(self, nid):
        cur = nid
        seen = 0
        while seen < 60:
            if self.nodes[cur]['type'] in COMP_TYPES:
                return cur
            if self.nodes[cur]['parent'] == -1:
                return None
            cur = self.nodes[cur]['parent']
            seen += 1
        return None

    def last_descendant(self, nid):
        """maior indice da subarvore -- usado para inserir filhos em sequencia"""
        return max(self.subtree(nid))

    # ------------------------------------------------------------- rebuild

    def rebuild(self, plan):
        """
        plan: lista ordenada de itens, cada um
              ('old', id_antigo)  ou  ('new', dict_do_no)
        Reconstroi compobj nessa ordem e remapeia TODAS as referencias a no.
        Linhas que apontam para nos ausentes do plano sao descartadas.
        """
        old_nodes = self.nodes
        mapping = {}
        new_nodes = []
        for new_id, item in enumerate(plan):
            if item[0] == 'old':
                mapping[item[1]] = new_id
                new_nodes.append(dict(old_nodes[item[1]]))
            else:
                new_nodes.append(dict(item[1]))

        def mp(x):
            return mapping.get(int(x))

        for n in new_nodes:
            # '_abs': o pai ja e um ID final (usado ao colar subarvores,
            # em que um no novo aponta para outro no novo)
            if n.pop('_abs', False) or n['parent'] == -1:
                continue
            tgt = mapping.get(n['parent'])
            n['parent'] = -1 if tgt is None else tgt
        self.nodes = new_nodes

        self.settings = [
            [mp(r[0]), r[1], str(mp(r[2])) if r[1] in SETTING_NODE_KEYS else r[2]]
            for r in self.settings
            if mp(r[0]) is not None
            and (r[1] not in SETTING_NODE_KEYS or mp(r[2]) is not None)
        ]

        for name in SIMPLE_FILES:
            out = []
            for r in self.rows[name]:
                t = mp(r[0])
                if t is None:
                    continue
                out.append([str(t)] + r[1:])
            self.rows[name] = out

        self.advancement = [
            [mp(r[0]), r[1], mp(r[2]), r[3]] for r in self.advancement
            if mp(r[0]) is not None and mp(r[2]) is not None
        ]

        out = []
        for r in self.tasks:
            if mp(r[0]) is None:
                continue
            r = list(r)
            r[0] = str(mp(r[0]))
            ok = True
            for i in TASK_NODE_FIELDS.get(r[2], []):
                t = mp(r[i])
                if t is None:
                    ok = False
                    break
                r[i] = str(t)
            if ok:
                out.append(r)
        self.tasks = out
        return mapping

    # ------------------------------------------------------------ mutacoes

    def add_node(self, parent, ntype, short, desc):
        self.snapshot()
        pos = self.last_descendant(parent) + 1
        plan = [('old', i) for i in range(pos)]
        plan.append(('new', {'type': ntype, 'short': short,
                             'desc': desc, 'parent': parent}))
        plan += [('old', i) for i in range(pos, len(self.nodes))]
        # o pai e um id antigo; como pos > parent, o mapeamento o preserva
        self.rebuild(plan)
        return pos

    def delete_node(self, nid):
        self.snapshot()
        parent = self.nodes[nid]['parent']
        dead = set(self.subtree(nid))
        mapping = self.rebuild([('old', i) for i in range(len(self.nodes))
                                if i not in dead])
        return {'removed': len(dead), 'parent': mapping.get(parent)}

    def move_node(self, nid, delta):
        """reordena um no entre seus irmaos, levando a subarvore junto"""
        parent = self.nodes[nid]['parent']
        sibs = self.children(parent)
        idx = sibs.index(nid)
        tgt = idx + delta
        if tgt < 0 or tgt >= len(sibs):
            return False
        self.snapshot()
        blocks = {s: self.subtree(s) for s in sibs}
        order = sibs[:]
        order[idx], order[tgt] = order[tgt], order[idx]
        moving = set()
        for s in sibs:
            moving |= set(blocks[s])
        first = min(moving)
        plan = [('old', i) for i in range(len(self.nodes))
                if i < first and i not in moving]
        for s in order:
            plan += [('old', i) for i in sorted(blocks[s])]
        plan += [('old', i) for i in range(len(self.nodes))
                 if i > first and i not in moving]
        mapping = self.rebuild(plan)
        return mapping.get(nid)

    def update_node(self, nid, short=None, desc=None):
        self.snapshot()
        if short is not None:
            self.nodes[nid]['short'] = short
        if desc is not None:
            self.nodes[nid]['desc'] = desc if desc.strip() else ' '

    # --------------------------------------------------------- copiar/colar

    def export_subtree(self, nid):
        """Serializa um no e tudo abaixo dele, com todas as linhas associadas.
        Referencias internas viram indices relativos; externas guardam o ID
        original e sao resolvidas na hora de colar (ou descartadas se o alvo
        nao existir mais)."""
        ids = self.subtree(nid)
        idset = set(ids)
        index = {old: i for i, old in enumerate(ids)}

        nodes = []
        for old in ids:
            n = self.nodes[old]
            nodes.append({'type': n['type'], 'short': n['short'], 'desc': n['desc'],
                          'parent': -1 if old == nid else index[n['parent']]})

        settings = []
        for r in self.settings:
            if r[0] not in idset:
                continue
            key, val = r[1], r[2]
            if key in SETTING_NODE_KEYS:
                try:
                    v = int(val)
                    val = ('i', index[v]) if v in idset else ('e', v)
                except ValueError:
                    pass
            settings.append([index[r[0]], key, val])

        tasks = []
        for r in self.tasks:
            if int(r[0]) not in idset:
                continue
            row = list(r)
            row[0] = index[int(r[0])]
            fields = []
            for i in TASK_NODE_FIELDS.get(r[2], []):
                v = int(r[i])
                fields.append((i, index[v] if v in idset else None, v))
            tasks.append({'row': row, 'fields': fields})

        return {
            'root_type': self.nodes[nid]['type'],
            'nodes': nodes,
            'settings': settings,
            'tasks': tasks,
            'advancement': [[index[r[0]], r[1], index[r[2]], r[3]]
                            for r in self.advancement
                            if r[0] in idset and r[2] in idset],
            'rows': {name: [[index[int(r[0])]] + list(r[1:])
                            for r in self.rows[name] if int(r[0]) in idset]
                     for name in SIMPLE_FILES},
        }

    # onde cada tipo de no pode ser colado
    PASTE_TARGETS = {3: CONTAINER_TYPES, 6: CONTAINER_TYPES,
                     4: COMP_TYPES, 5: (4, 5, 6), 2: (0, 1), 1: (0,)}

    def import_subtree(self, parent, data):
        """Cola uma subarvore exportada como filho de `parent`."""
        rt = data['nodes'][0]['type']
        pt = self.nodes[parent]['type']
        if rt not in self.PASTE_TARGETS or pt not in self.PASTE_TARGETS[rt]:
            raise ValueError('nao da para colar %s dentro de %s'
                             % (NODE_TYPES.get(rt, rt), NODE_TYPES.get(pt, pt)))

        self.snapshot()
        pos = self.last_descendant(parent) + 1
        new_nodes = []
        for n in data['nodes']:
            new_nodes.append({'type': n['type'], 'short': n['short'],
                              'desc': n['desc'], '_abs': True,
                              'parent': parent if n['parent'] == -1
                              else pos + n['parent']})
        plan = [('old', i) for i in range(pos)] \
             + [('new', d) for d in new_nodes] \
             + [('old', i) for i in range(pos, len(self.nodes))]
        mapping = self.rebuild(plan)

        clip = lambda idx: pos + idx

        for idx, key, val in data['settings']:
            if isinstance(val, tuple):
                tgt = clip(val[1]) if val[0] == 'i' else mapping.get(val[1])
                if tgt is None:
                    continue
                val = str(tgt)
            self.settings.append([clip(idx), key, val])

        for name, rows in data['rows'].items():
            for r in rows:
                self.rows[name].append([str(clip(r[0]))] + list(r[1:]))

        for a, rk, b, sl in data['advancement']:
            self.advancement.append([clip(a), rk, clip(b), sl])

        for t in data['tasks']:
            row = list(t['row'])
            row[0] = str(clip(row[0]))
            ok = True
            for i, internal, raw in t['fields']:
                if internal is not None:
                    row[i] = str(clip(internal))
                else:
                    tgt = mapping.get(raw)
                    if tgt is None:
                        ok = False
                        break
                    row[i] = str(tgt)
            if ok:
                self.tasks.append(row)

        return clip(0)

    def set_slots(self, gid, count):
        """define quantas vagas (linhas de standings) o grupo tem"""
        self.snapshot()
        self.rows['standings'] = [r for r in self.rows['standings']
                                  if int(r[0]) != gid]
        anchor = None
        for i, r in enumerate(self.rows['standings']):
            if int(r[0]) < gid:
                anchor = i
        ins = (anchor + 1) if anchor is not None else len(self.rows['standings'])
        new = [[str(gid), str(k)] for k in range(count)]
        self.rows['standings'][ins:ins] = new

    def set_settings(self, nid, rows):
        self.snapshot()
        keep = [r for r in self.settings if r[0] != nid]
        first = next((i for i, r in enumerate(self.settings) if r[0] == nid), None)
        block = [[nid, k, v] for k, v in rows]
        if first is None:
            self.settings = keep + block
        else:
            before = [r for r in self.settings[:first] if r[0] != nid]
            after = [r for r in self.settings[first:] if r[0] != nid]
            self.settings = before + block + after

    def set_file_rows(self, name, nid, rows):
        """substitui todas as linhas de um dos arquivos simples para um no"""
        self.snapshot()
        cur = self.rows[name]
        first = next((i for i, r in enumerate(cur) if int(r[0]) == nid), None)
        block = [[str(nid)] + [str(x) for x in r] for r in rows]
        if first is None:
            anchor = None
            for i, r in enumerate(cur):
                if int(r[0]) < nid:
                    anchor = i
            ins = (anchor + 1) if anchor is not None else len(cur)
            cur[ins:ins] = block
        else:
            before = [r for r in cur[:first] if int(r[0]) != nid]
            after = [r for r in cur[first:] if int(r[0]) != nid]
            self.rows[name] = before + block + after

    def set_advancement(self, dst, rows):
        """rows = [(src, rank, slot)] -- todas as entradas que chegam em dst"""
        self.snapshot()
        first = next((i for i, r in enumerate(self.advancement) if r[2] == dst), None)
        block = [[int(s), int(rk), dst, int(sl)] for s, rk, sl in rows]
        keep_before = [r for r in self.advancement[:first or 0] if r[2] != dst]
        keep_after = [r for r in self.advancement[(first or 0):] if r[2] != dst]
        self.advancement = keep_before + block + keep_after

    def set_tasks(self, nid, rows):
        self.snapshot()
        first = next((i for i, r in enumerate(self.tasks) if int(r[0]) == nid), None)
        block = [[str(nid)] + [str(x) for x in r] for r in rows]
        before = [r for r in self.tasks[:first or 0] if int(r[0]) != nid]
        after = [r for r in self.tasks[(first or 0):] if int(r[0]) != nid]
        self.tasks = before + block + after

    # ------------------------------------------------------------- consultas

    def settings_of(self, nid):
        return [[r[1], r[2]] for r in self.settings if r[0] == nid]

    def setting(self, nid, key, default=None):
        for r in self.settings:
            if r[0] == nid and r[1] == key:
                return r[2]
        return default

    def slots_of(self, nid):
        return sum(1 for r in self.rows['standings'] if int(r[0]) == nid)

    def rows_of(self, name, nid):
        return [r[1:] for r in self.rows[name] if int(r[0]) == nid]

    def tasks_of(self, nid):
        return [r[1:] for r in self.tasks if int(r[0]) == nid]

    def adv_in(self, nid):
        return [[r[0], r[1], r[3]] for r in self.advancement if r[2] == nid]

    def adv_out(self, nid):
        return [[r[1], r[2], r[3]] for r in self.advancement if r[0] == nid]

    def teams_of(self, nid):
        return [int(v) for k, v in self.settings_of(nid) if k == 'info_special_team_id']

    def impact(self, nid):
        """quantas linhas cada arquivo perde ao remover a subarvore de nid"""
        dead = set(self.subtree(nid))
        out = {
            'nos': len(dead),
            'compids': sum(1 for i in dead if self.nodes[i]['type'] in COMP_TYPES),
            'settings': sum(1 for r in self.settings if r[0] in dead),
            'advancement': sum(1 for r in self.advancement
                               if r[0] in dead or r[2] in dead),
            'tasks': sum(1 for r in self.tasks if int(r[0]) in dead),
        }
        for name in SIMPLE_FILES:
            out[name] = sum(1 for r in self.rows[name] if int(r[0]) in dead)
        return out

    # ------------------------------------------------------------- validacao

    def validate(self):
        errors, warnings = [], []

        def E(nid, msg):
            errors.append({'node': nid, 'msg': msg})

        def W(nid, msg):
            warnings.append({'node': nid, 'msg': msg})

        n_nodes = len(self.nodes)
        for i, n in enumerate(self.nodes):
            p = n['parent']
            if p == -1:
                continue
            if p >= n_nodes or p < -1:
                E(i, 'pai %d nao existe' % p)
                continue
            pt = self.nodes[p]['type']
            if n['type'] == 4 and pt not in COMP_TYPES:
                E(i, 'estagio pendurado num %s' % NODE_TYPES.get(pt, pt))
            # grupo normalmente fica sob um estagio, mas a EA tambem pendura
            # grupo direto numa competicao especial (C90001..C90010) e dentro
            # de outro grupo. Ambos os casos sao validos.
            if n['type'] == 5 and pt not in (4, 5, 6):
                E(i, 'grupo pendurado num %s' % NODE_TYPES.get(pt, pt))
            if n['type'] == 3 and pt not in CONTAINER_TYPES:
                W(i, 'competicao pendurada num %s' % NODE_TYPES.get(pt, pt))

        slots = collections.Counter()
        pos = collections.defaultdict(list)
        for r in self.rows['standings']:
            a = int(r[0])
            if a >= n_nodes:
                E(None, 'standings aponta para o no inexistente %d' % a)
                continue
            if self.nodes[a]['type'] != 5:
                E(a, 'standings num no que nao e grupo')
            slots[a] += 1
            pos[a].append(int(r[1]))
        for g, ps in pos.items():
            if sorted(ps) != list(range(len(ps))):
                E(g, 'posicoes de standings fora de sequencia: %s' % sorted(ps))

        for r in self.settings:
            if r[0] >= n_nodes:
                E(None, 'settings aponta para o no inexistente %d' % r[0])

        fill = collections.Counter()
        for r in self.advancement:
            a, rank, b, slot = r
            if a >= n_nodes or b >= n_nodes:
                E(None, 'advancement aponta para no inexistente (%s)' % (r,))
                continue
            if self.nodes[a]['type'] != 5 or self.nodes[b]['type'] != 5:
                E(a, 'advancement precisa ligar grupo a grupo')
            if a in slots and rank > slots[a]:
                E(a, 'avanca da posicao %d mas o grupo so tem %d vagas' % (rank, slots[a]))
            if b in slots and not (1 <= slot <= slots[b]):
                E(b, 'chega na vaga %d mas o grupo so tem %d' % (slot, slots[b]))
            if rank != 0:
                fill[(b, slot)] += 1
        for (b, slot), c in fill.items():
            if c > 1:
                E(b, 'vaga %d preenchida %d vezes' % (slot, c))

        for r in self.rows['schedule']:
            a = int(r[0])
            if a >= n_nodes:
                E(None, 'schedule aponta para o no inexistente %d' % a)
                continue
            if self.nodes[a]['type'] not in (4, 5):
                E(a, 'schedule num no que nao e estagio nem grupo')
            if len(r) >= 5 and int(r[3]) > int(r[4]):
                E(a, 'schedule com minimo de jogos maior que o maximo')

        for r in self.tasks:
            a = int(r[0])
            if a >= n_nodes:
                E(None, 'tasks aponta para o no inexistente %d' % a)
                continue
            if self.nodes[a]['type'] not in COMP_TYPES:
                E(a, 'task precisa comecar numa competicao')
            if r[2] not in TASK_NODE_FIELDS:
                W(a, 'verbo de task desconhecido: %s' % r[2])
                continue
            for i in TASK_NODE_FIELDS[r[2]]:
                if int(r[i]) >= n_nodes:
                    E(a, 'task aponta para o no inexistente %s' % r[i])
            if r[2] == 'FillFromSpecialTeams':
                tgt, want = int(r[3]), int(r[4])
                pool = len(self.teams_of(tgt))
                if pool < want:
                    E(tgt, 'a task pede %d times mas so ha %d na lista' % (want, pool))
                if slots.get(tgt, 0) < want:
                    E(tgt, 'a task pede %d times mas o grupo tem %d vagas'
                      % (want, slots.get(tgt, 0)))

        for name in ('objectives', 'initteams', 'weather'):
            for r in self.rows[name]:
                if int(r[0]) >= n_nodes:
                    E(None, '%s aponta para o no inexistente %s' % (name, r[0]))

        # 'slots' ja conta as linhas de standings de cada grupo (= slots_of) e 'com_filho' troca
        # children(i) por um conjunto: mesmo resultado, sem varrer os arquivos de novo para cada no
        com_filho = set(n['parent'] for n in self.nodes)
        for i, n in enumerate(self.nodes):
            if n['type'] == -1:
                continue                       # linha reservada, so ocupa posicao
            if n['type'] == 5 and slots.get(i, 0) == 0:
                W(i, 'grupo sem nenhuma vaga em standings')
            if n['type'] == 4 and i not in com_filho:
                W(i, 'estagio sem grupos')
            if n['type'] in COMP_TYPES and i not in com_filho:
                W(i, 'competicao sem estagios')

        return {'errors': errors, 'warnings': warnings}


# ------------------------------------------------------------------ nomes

class Names(object):
    """Nomes de times e competicoes lidos da planilha, para a interface."""

    def __init__(self, path):
        self.teams = {}
        self.comps = {}
        self.planilha = None
        book = None
        for fn in sorted(os.listdir(path)):
            if fn.lower().endswith('.xlsx') and not fn.startswith('~'):
                book = os.path.join(path, fn)
                break
        if not book:
            return
        self.planilha = os.path.basename(book)
        try:
            import openpyxl
        except ImportError:
            openpyxl = None
        if openpyxl is None:
            # sem openpyxl: leitor mínimo de .xlsx com a biblioteca padrão (só valores)
            try:
                for sheet, rows in _xlsx_planilhas(book, ('time', 'competi')).items():
                    into = self.teams if sheet.lower().startswith('time') else self.comps
                    for row in rows:
                        self._collect(row, into)
            except Exception:
                pass
            return
        try:
            wb = openpyxl.load_workbook(book, data_only=True, read_only=True)
        except Exception:
            return
        for sheet in wb.sheetnames:
            low = sheet.lower()
            if low.startswith('time'):
                for row in wb[sheet].iter_rows(values_only=True):
                    self._collect(row, self.teams)
            elif low.startswith('competi'):
                for row in wb[sheet].iter_rows(values_only=True):
                    self._collect(row, self.comps)
        try:
            wb.close()
        except Exception:
            pass

    @staticmethod
    def _collect(row, into):
        vals = [c for c in row if c is not None]
        num = next((c for c in vals if isinstance(c, (int, float))
                    and float(c) == int(c) and 0 < int(c) < 10 ** 7), None)
        txt = next((str(c).strip() for c in vals
                    if isinstance(c, str) and str(c).strip()), None)
        if num is not None and txt:
            into.setdefault(int(num), txt)

    def team(self, tid):
        return self.teams.get(int(tid))

    def comp(self, aid):
        return self.comps.get(int(aid))


def _xlsx_planilhas(book, prefixos):
    """Linhas (tuplas de valores) das planilhas cujo nome começa com um dos prefixos. Só biblioteca padrão."""
    import zipfile
    import xml.etree.ElementTree as ET
    m = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
    r_id = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'
    out = {}
    with zipfile.ZipFile(book) as z:
        nomes = set(z.namelist())
        shared = []
        if 'xl/sharedStrings.xml' in nomes:
            for si in ET.fromstring(z.read('xl/sharedStrings.xml')).iter(m + 'si'):
                shared.append(''.join(t.text or '' for t in si.iter(m + 't')))
        alvo = {rel.get('Id'): rel.get('Target', '') for rel in ET.fromstring(z.read('xl/_rels/workbook.xml.rels'))}
        for sh in ET.fromstring(z.read('xl/workbook.xml')).iter(m + 'sheet'):
            nome = sh.get('name') or ''
            if not nome.lower().startswith(prefixos):
                continue
            t = alvo.get(sh.get(r_id), '').lstrip('/')
            arq = t if t.startswith('xl/') else 'xl/' + t
            if arq not in nomes:
                continue
            linhas = []
            for row in ET.fromstring(z.read(arq)).iter(m + 'row'):
                vals = []
                for c in row.iter(m + 'c'):
                    tipo, v = c.get('t'), c.find(m + 'v')
                    if tipo == 's' and v is not None:
                        vals.append(shared[int(v.text)])
                    elif tipo == 'inlineStr':
                        vals.append(''.join(x.text or '' for x in c.iter(m + 't')))
                    elif v is None or v.text is None:
                        vals.append(None)
                    elif tipo in ('str', 'e'):
                        vals.append(v.text)
                    elif tipo == 'b':
                        vals.append(v.text == '1')
                    else:
                        try:
                            f = float(v.text)
                            vals.append(int(f) if f.is_integer() else f)
                        except ValueError:
                            vals.append(v.text)
                linhas.append(tuple(vals))
            out[nome] = linhas
    return out


_RAW_KEY = re.compile(r'^[A-Za-z]+(_[A-Za-z0-9]+)+$')


def label_for(cd, names, nid):
    """rotulo amigavel de um no, usando asset_id quando existir (e o idioma do jogo, se houver um importado)"""
    n = cd.nodes[nid]
    texto = getattr(names, 'texto', None) if names else None
    if n['type'] == -1:
        return 'reservado'
    if n['type'] in CONTAINER_TYPES:
        d = n['desc'].strip()
        # NationName_54 e chave de traducao: com o idioma vira "Brasil"; sem ele, a sigla (BRAZ) diz mais
        if d and _RAW_KEY.match(d):
            return (texto(d) if texto else None) or n['short']
        return d or n['short']
    if n['type'] in COMP_TYPES:
        aid = cd.setting(nid, 'asset_id')
        if aid and names:
            nm = names.comp(aid)
            if nm:
                return nm
        m = re.match(r'^C(\d+)$', n['short'])
        if m and names:
            nm = names.comp(m.group(1))
            if nm:
                return nm
        d = n['desc'].strip()
        if texto and d:
            nm = texto(d)
            if nm:
                return nm
        return n['short']
    d = n['desc'].strip()
    if texto and d and _RAW_KEY.match(d):
        d = texto(d) or d
    if n['type'] == 5:
        return '%s%s' % (n['short'], (' - ' + d) if d else '')
    return d if d else n['short']
